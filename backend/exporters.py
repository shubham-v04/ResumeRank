"""
Turns the ranked results list into downloadable files.
Kept separate from main.py so the export format can be changed
without touching the API routes.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from fpdf import FPDF

# Brand colors, matching the website's teal theme
ACCENT_RGB = (31, 122, 92)       # #1F7A5C
ACCENT_DARK_RGB = (18, 42, 46)    # #122A2E
ROW_ALT_RGB = (228, 242, 236)     # light teal, #E4F2EC


def _skill_summary(row):
    if row.get("skill_match_percent") is None:
        return "N/A"
    total = len(row["skills_matched"]) + len(row["skills_missing"])
    return f"{row['skill_match_percent']}% ({len(row['skills_matched'])}/{total})"


def _experience_summary(row):
    if row.get("years_experience") is None:
        return "Unknown"
    if row.get("meets_experience") is None:
        return f"{row['years_experience']} yrs"
    flag = "OK" if row["meets_experience"] else "Below min"
    return f"{row['years_experience']} yrs ({flag})"


def generate_excel(results: list[dict]) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ranked Resumes"

    headers = ["Rank", "Name", "Email", "Match Score (%)", "Skill Match", "Experience", "Filename"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F7A5C", end_color="1F7A5C", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    thin_border = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )
    alt_fill = PatternFill(start_color="E4F2EC", end_color="E4F2EC", fill_type="solid")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for i, row in enumerate(results, start=1):
        ws.append([
            i, row["name"], row["email"], row["score"],
            _skill_summary(row), _experience_summary(row), row["filename"],
        ])
        excel_row = ws[i + 1]
        for cell in excel_row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if i % 2 == 0:
                cell.fill = alt_fill

    widths = [6, 22, 28, 16, 20, 20, 34]
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 20

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_pdf(results: list[dict]) -> io.BytesIO:
    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.add_page()

    # Title bar
    pdf.set_fill_color(*ACCENT_DARK_RGB)
    pdf.rect(0, 0, 297, 18, style="F")
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 16)
    pdf.set_xy(10, 4)
    pdf.cell(0, 10, "ResumeRank - Ranked Results", ln=True)
    pdf.ln(6)

    # Column widths, sized so the filename fits without clipping
    col_widths = [10, 32, 45, 20, 32, 28, 65]
    headers = ["Rank", "Name", "Email", "Score", "Skill Match", "Experience", "Filename"]

    def draw_header_row():
        pdf.set_font("Helvetica", "B", 9)
        pdf.set_fill_color(*ACCENT_RGB)
        pdf.set_text_color(255, 255, 255)
        pdf.set_draw_color(255, 255, 255)
        for header, width in zip(headers, col_widths):
            pdf.cell(width, 9, header, border=1, fill=True, align="C")
        pdf.ln()

    draw_header_row()

    pdf.set_font("Helvetica", "", 8)
    pdf.set_draw_color(210, 210, 210)

    for i, row in enumerate(results, start=1):
        # Alternate row shading
        if i % 2 == 0:
            pdf.set_fill_color(*ROW_ALT_RGB)
            fill = True
        else:
            pdf.set_fill_color(255, 255, 255)
            fill = True

        pdf.set_text_color(30, 30, 30)

        values = [
            str(i), row["name"], row["email"], f"{row['score']}%",
            _skill_summary(row), _experience_summary(row), row["filename"],
        ]

        # Truncate only if a value still wouldn't fit even at this width
        max_chars = [6, 20, 28, 8, 20, 18, 45]
        for value, width, limit in zip(values, col_widths, max_chars):
            text = value if len(value) <= limit else value[: limit - 1] + "."
            pdf.cell(width, 8, text, border=1, fill=fill, align="L")
        pdf.ln()

        # Re-draw header every ~25 rows so long lists stay readable across page breaks
        if i % 25 == 0 and i != len(results):
            pdf.add_page()
            draw_header_row()
            pdf.set_font("Helvetica", "", 8)

    buffer = io.BytesIO()
    pdf.output(buffer)
    buffer.seek(0)
    return buffer